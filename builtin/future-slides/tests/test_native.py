from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
import native_deck as native


class NativeDeckTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="native-deck-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.manifest = self.root / "deck.json"
        self.spec = json.loads((SKILL / "assets" / "example-native.json").read_text(encoding="utf-8"))
        self.save()

    def save(self):
        self.manifest.write_text(json.dumps(self.spec, ensure_ascii=False), encoding="utf-8")

    def add_image(self):
        path = self.root / "original figure.png"
        Image.new("RGB", (240, 120), "white").save(path)
        element = {"id": "original", "kind": "image", "x": 8.3, "y": 5.2, "w": 3.5, "h": 1.1,
                   "path": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                   "alt": "Synthetic blank image; no measurements"}
        self.spec["slides"][1]["elements"].append(element)
        self.save()
        return path, element

    def add_table(self):
        self.spec["slides"].append({"id": "table", "title": "Synthetic table", "elements": [
            {"id": "values", "kind": "table", "x": .5, "y": 1, "w": 8, "h": 3,
             "rows": [["Condition", "Synthetic response"], ["A", "1.2"], ["B", "0.7"]]}]})
        self.save()

    def add_category_chart(self):
        self.spec["slides"].append({"id": "categories", "title": "Synthetic chart", "elements": [
            {"id": "chart", "kind": "chart", "x": .5, "y": 1, "w": 8, "h": 5,
             "chart_type": "column", "categories": ["A", "B"],
             "series": [{"name": "Synthetic", "values": [1.25, -.5]}], "source_ids": ["demo"]}]})
        self.save()

    def test_example_build_check_native_objects_and_true_bindings(self):
        report = native.build(self.manifest, "talk.pptx")
        checked = native.check(self.manifest, "talk.pptx")
        self.assertEqual(checked["structural_check"], "passed")
        self.assertEqual([s["id"] for s in checked["slides"]], ["design", "observations"])
        self.assertEqual(report["visual_review"], "not_performed")
        self.assertEqual(report["native_render"], "not_performed")
        self.assertEqual(report["content_review"], "not_performed")
        with zipfile.ZipFile(self.root / "talk.pptx") as z:
            self.assertFalse(any(n.startswith("ppt/media/") for n in z.namelist()))
            self.assertEqual(sum(n.startswith("ppt/embeddings/") and n.endswith(".xlsx") for n in z.namelist()), 1)
        prs = Presentation(self.root / "talk.pptx")
        shapes = {s.name: s for s in prs.slides[0].shapes}
        start = shapes["link-a"]._element.xpath(".//a:stCxn")[0]
        end = shapes["link-a"]._element.xpath(".//a:endCxn")[0]
        self.assertEqual(start.get("id"), str(shapes["question"].shape_id))
        self.assertEqual(end.get("id"), str(shapes["measure"].shape_id))
        self.assertEqual(start.get("idx"), "3")
        chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
        self.assertEqual(native.chart_values(chart.series[0], "xVal"), [0, 2, 10])

    def test_real_text_shape_table_chart_edit_save_reopen(self):
        self.add_table(); self.add_category_chart()
        prs = native.create(self.spec, self.root)
        shapes = {s.name: s for s in prs.slides[0].shapes}
        shapes["title"].text = "修改标题 — revised title"
        shapes["question"].fill.fore_color.rgb = RGBColor.from_string("123456")
        prs.slides[2].shapes[0].table.cell(1, 1).text = "2.5"
        data = CategoryChartData(); data.categories = ["A", "B"]; data.add_series("Revised", [3, 4])
        prs.slides[3].shapes[0].chart.replace_data(data)
        stream = BytesIO(); prs.save(stream); stream.seek(0); reopened = Presentation(stream)
        new_shapes = {s.name: s for s in reopened.slides[0].shapes}
        self.assertEqual(new_shapes["title"].text, "修改标题 — revised title")
        self.assertEqual(str(new_shapes["question"].fill.fore_color.rgb), "123456")
        self.assertEqual(reopened.slides[2].shapes[0].table.cell(1, 1).text, "2.5")
        self.assertEqual(list(reopened.slides[3].shapes[0].chart.series[0].values), [3, 4])

    def test_image_bytes_aspect_ratio_and_no_crop(self):
        _, element = self.add_image()
        native.build(self.manifest, "images.pptx")
        native.check(self.manifest, "images.pptx")
        prs = Presentation(self.root / "images.pptx")
        pic = next(s for s in prs.slides[1].shapes if s.name == "original")
        self.assertAlmostEqual(pic.width / pic.height, 2)
        self.assertEqual(hashlib.sha256(pic.image.blob).hexdigest(), element["sha256"])
        self.assertEqual(pic.crop_top, 0)
        self.assertEqual(pic._element.nvPicPr.cNvPr.get("descr"), element["alt"])

    def test_all_category_types_and_table_reopen(self):
        self.add_table(); self.add_category_chart()
        for chart_type in ("column", "bar", "line"):
            with self.subTest(chart_type=chart_type):
                e = self.spec["slides"][-1]["elements"][0]
                e.update(chart_type=chart_type, x_title="Horizontal", y_title="Vertical")
                self.save()
                native.build(self.manifest, chart_type + ".pptx")
                native.check(self.manifest, chart_type + ".pptx")
                prs = Presentation(self.root / (chart_type + ".pptx"))
                chart = prs.slides[-1].shapes[0].chart
                horizontal = chart.value_axis if chart_type == "bar" else chart.category_axis
                self.assertEqual(horizontal.axis_title.text_frame.text, "Horizontal")

    def test_changed_source_stale_release_and_no_overwrite(self):
        native.build(self.manifest, "talk.pptx")
        before = (self.root / "talk.pptx").read_bytes()
        with self.assertRaisesRegex(ValueError, "already exists"):
            native.build(self.manifest, "talk.pptx")
        self.assertEqual(before, (self.root / "talk.pptx").read_bytes())
        self.spec["slides"][0]["title"] = "New revision"; self.save()
        with self.assertRaisesRegex(ValueError, "Stale source"):
            native.check(self.manifest, "talk.pptx")
        native.build(self.manifest, "talk-r2.pptx")
        self.assertEqual(before, (self.root / "talk.pptx").read_bytes())

    def test_manual_pptx_edits_detected(self):
        native.build(self.manifest, "talk.pptx")
        path = self.root / "talk.pptx"
        prs = Presentation(path); prs.slides[0].shapes[0].text = "Manual edit"; prs.save(path)
        with self.assertRaisesRegex(ValueError, "PPTX changed"):
            native.check(self.manifest, "talk.pptx")

    def test_image_change_rejected(self):
        path, _ = self.add_image()
        native.build(self.manifest, "talk.pptx")
        Image.new("RGB", (240, 120), "black").save(path)
        with self.assertRaisesRegex(ValueError, "image hash mismatch"):
            native.check(self.manifest, "talk.pptx")

    def test_bad_geometry_ids_refs_and_types_rejected(self):
        mutations = [
            lambda s: s["slides"][0]["elements"][0].update(x=-1),
            lambda s: s["slides"][0]["elements"][0].update(w=0),
            lambda s: s["slides"][0]["elements"][0].update(h=100),
            lambda s: s["slides"][0]["elements"][0].update(size=float("nan")),
            lambda s: s["slides"][0]["elements"][0].update(x=True),
            lambda s: s["slides"][0]["elements"][0].update(kind="unsupported"),
            lambda s: s["slides"][0]["elements"][0].update(error_bars="silently ignored?"),
            lambda s: s["slides"][0]["elements"][0].update(source_ids=["unknown"]),
            lambda s: s["slides"][0]["elements"][1].update(id="title"),
            lambda s: s["slides"][1].update(id="design"),
            lambda s: s["slides"][0]["elements"][0].update(text="bad\x00text"),
        ]
        for mutate in mutations:
            with self.subTest(mutation=mutations.index(mutate)):
                spec = deepcopy(self.spec); mutate(spec)
                with self.assertRaises(ValueError):
                    native.validate(spec, self.root)

    def test_chart_invalid_data_never_coerced_or_invented(self):
        for points in ([[0, None]], [[0, float("inf")]], [[0]], []):
            spec = deepcopy(self.spec)
            spec["slides"][1]["elements"][1]["series"][0]["points"] = points
            with self.subTest(points=points), self.assertRaises(ValueError):
                native.validate(spec, self.root)
        self.add_category_chart()
        self.spec["slides"][-1]["elements"][0]["series"][0]["values"] = [1]
        with self.assertRaisesRegex(ValueError, "length mismatch"):
            native.validate(self.spec, self.root)

    def test_paths_are_portable_and_confined(self):
        for value in ("../escape.pptx", "/absolute.pptx", "C:/escape.pptx", "C:escape.pptx",
                      "\\\\server\\deck.pptx", "nested\\deck.pptx"):
            with self.subTest(path=value), self.assertRaises(ValueError):
                native.output_paths(self.root, value)
        self.assertEqual(native.output_paths(self.root, "nested/talk.pptx")[0], self.root / "nested" / "talk.pptx")
        with self.assertRaises(ValueError):
            native.output_paths(self.root, "deck.json")

    def test_symlink_escape_rejected_where_supported(self):
        with tempfile.TemporaryDirectory() as outside:
            link = self.root / "outside"
            try:
                link.symlink_to(outside, target_is_directory=True)
            except (OSError, NotImplementedError):
                self.skipTest("Directory symlinks unavailable for current account")
            with self.assertRaisesRegex(ValueError, "escapes"):
                native.output_paths(self.root, "outside/deck.pptx")

    def test_scene_connector_wrong_target_rejected(self):
        connector = next(e for e in self.spec["slides"][0]["elements"] if e["kind"] == "connector")
        connector["target"]["id"] = "title"
        with self.assertRaisesRegex(ValueError, "native shapes"):
            native.validate(self.spec, self.root)

    def test_saved_connector_and_chart_corruption_rejected(self):
        prs = native.create(self.spec, self.root)
        link = next(s for s in prs.slides[0].shapes if s.name == "link-a")
        node = link._element.xpath(".//a:endCxn")[0]; node.set("id", "99999")
        with self.assertRaisesRegex(ValueError, "detached"):
            native.inspect(prs, self.spec, self.root)
        prs = native.create(self.spec, self.root)
        plot = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
        plot.series[0]._element.xpath("./c:xVal/c:numRef/c:numCache/c:pt/c:v")[0].text = "7"
        with self.assertRaisesRegex(ValueError, "X values changed"):
            native.inspect(prs, self.spec, self.root)

    def test_shape_variants_elbow_and_unarrowed_connector(self):
        elements = self.spec["slides"][0]["elements"]
        next(e for e in elements if e["id"] == "question")["shape"] = "ellipse"
        next(e for e in elements if e["id"] == "measure")["shape"] = "round_rect"
        link = next(e for e in elements if e["id"] == "link-a")
        link.update(routing="elbow", arrow=False)
        self.save()
        native.build(self.manifest, "variants.pptx")
        native.check(self.manifest, "variants.pptx")
        prs = Presentation(self.root / "variants.pptx")
        shapes = {s.name: s for s in prs.slides[0].shapes}
        self.assertFalse(shapes["link-a"]._element.xpath(".//a:tailEnd"))
        head = shapes["link-b"]._element.xpath(".//a:tailEnd")[0]
        head.set("type", "none")
        with self.assertRaisesRegex(ValueError, "Arrow mismatch"):
            native.inspect(prs, self.spec, self.root)

    def test_language_canvas_notes_and_order_are_not_fixed(self):
        self.spec["language"] = "zh-CN"
        self.spec["title"] = "跨学科方法讨论"
        self.spec["width"] = 14; self.spec["height"] = 10.5
        self.spec["slides"].reverse()
        self.spec["slides"][0]["speaker_notes"] = "公开演讲者备注，不是内部审核记录。"
        self.spec["slides"][0]["elements"][0]["text"] = "研究问题 α 与不确定性"
        self.save()
        native.build(self.manifest, "中文报告.pptx")
        checked = native.check(self.manifest, "中文报告.pptx")
        self.assertEqual([s["id"] for s in checked["slides"]], ["observations", "design"])
        prs = Presentation(self.root / "中文报告.pptx")
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 4/3)
        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, "公开演讲者备注，不是内部审核记录。")

    def test_overlap_warns_without_fabricating_visual_pass(self):
        self.spec["slides"][0]["elements"][1].update(x=.55, y=.45)
        self.save()
        report = native.build(self.manifest, "overlap.pptx")
        self.assertTrue(any("overlap" in w for w in report["warnings"]))
        self.assertEqual(report["visual_review"], "not_performed")

    def test_cli_failure_nonzero_and_no_partial_deck(self):
        self.spec["slides"][0]["elements"][0]["x"] = -5; self.save()
        with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
            self.assertEqual(native.main(["build", str(self.manifest), "--output", "bad.pptx"]), 1)
        self.assertFalse((self.root / "bad.pptx").exists())
        self.assertFalse((self.root / "bad.build.json").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
