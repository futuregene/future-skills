#!/usr/bin/env python3
"""Local native PPTX authoring and hash-bound structural checks (not visual QA).

The explicit scene is the source of truth. No network, OCR, generation, rendering,
or implicit conversion of unsupported objects. Coordinates are in inches.
"""
import argparse
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path, PureWindowsPath
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

PORTS = {"top": 0, "left": 1, "bottom": 2, "right": 3}
CHARTS = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
          "line": XL_CHART_TYPE.LINE, "scatter": XL_CHART_TYPE.XY_SCATTER}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
          "ellipse": MSO_SHAPE.OVAL}
COMMON = {"id", "kind", "source_ids"}
BOX = {"x", "y", "w", "h"}
TEXT = {"text", "size", "color", "bold", "align"}
FIELDS = {
    "text": COMMON | BOX | TEXT,
    "shape": COMMON | BOX | TEXT | {"shape", "fill", "line"},
    "image": COMMON | BOX | {"path", "sha256", "alt"},
    "table": COMMON | BOX | {"rows", "size", "header"},
    "chart": COMMON | BOX | {"chart_type", "categories", "series", "x_title", "y_title"},
    "connector": COMMON | {"source", "target", "routing", "color", "arrow"},
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def local_path(root, value):
    require(isinstance(value, str) and value and "\\" not in value,
            "Use nonempty, forward-slash relative paths")
    require(not Path(value).is_absolute() and not PureWindowsPath(value).drive,
            "Absolute/drive paths are not allowed inside the deck")
    path = (root / value).resolve()
    require(path.is_relative_to(root), f"Path escapes deck directory: {value}")
    return path


def number(value, label, positive=False):
    require(type(value) in (int, float) and math.isfinite(value), f"{label}: finite number required")
    require(not positive or value > 0, f"{label}: must be positive")
    return value


def string(value, label):
    require(isinstance(value, str) and value.strip(), f"{label}: nonempty string required")
    require(not any(ord(c) < 32 and c not in "\n\t" for c in value), f"{label}: control character")
    return value


def color(value):
    require(isinstance(value, str) and re.fullmatch(r"[0-9A-Fa-f]{6}", value),
            f"Expected six-digit RGB color, got {value!r}")
    return RGBColor.from_string(value.upper())


def keys(obj, allowed, label):
    require(isinstance(obj, dict), f"{label}: object required")
    require(not (set(obj) - allowed), f"{label}: unsupported fields {sorted(set(obj) - allowed)}")


def safe_id(value):
    require(isinstance(value, str) and re.fullmatch(r"[A-Za-z0-9_-]+", value), "IDs must be safe strings")
    return value


def validate(spec, root):
    keys(spec, {"schema_version", "title", "language", "width", "height", "theme", "sources", "slides"}, "deck")
    require(type(spec.get("schema_version")) is int and spec["schema_version"] == 1, "schema_version must be 1")
    string(spec.get("title"), "title"); string(spec.get("language"), "language")
    width = number(spec.get("width"), "width", True)
    height = number(spec.get("height"), "height", True)
    require(width <= 56 and height <= 56, "Canvas exceeds 56-inch PPTX authoring limit")
    theme = spec.get("theme")
    keys(theme, {"font", "cjk_font", "background", "foreground", "accent"}, "theme")
    for field in ("font", "cjk_font"):
        string(theme.get(field), field)
    for field in ("background", "foreground", "accent"):
        color(theme.get(field))
    sources = spec.get("sources", [])
    require(isinstance(sources, list), "sources must be an array")
    source_ids = set()
    for src in sources:
        keys(src, {"id", "citation", "locator"}, "source")
        sid = safe_id(src.get("id"))
        require(sid not in source_ids, "Duplicate source ID")
        source_ids.add(sid)
        string(src.get("citation"), "citation"); string(src.get("locator"), "locator")
    slides = spec.get("slides")
    require(isinstance(slides, list) and slides, "slides must be a nonempty array in page order")
    slide_ids, warnings = set(), []
    for slide in slides:
        keys(slide, {"id", "title", "elements", "speaker_notes"}, "slide")
        sid = safe_id(slide.get("id"))
        require(sid not in slide_ids, "Duplicate slide ID")
        slide_ids.add(sid)
        string(slide.get("title"), "slide title")
        if "speaker_notes" in slide:
            string(slide["speaker_notes"], "speaker_notes")
        elements = slide.get("elements")
        require(isinstance(elements, list) and elements, f"{sid}: elements must be nonempty")
        by_id = {}
        for e in elements:
            require(isinstance(e, dict) and e.get("kind") in FIELDS, "Unsupported element kind")
            kind = e["kind"]
            keys(e, FIELDS[kind], kind)
            eid = safe_id(e.get("id"))
            require(eid not in by_id, f"{sid}: duplicate element ID {eid}")
            by_id[eid] = e
            refs = e.get("source_ids", [])
            require(isinstance(refs, list) and all(isinstance(r, str) and r in source_ids for r in refs),
                    f"{sid}/{eid}: unknown source reference")
            if kind != "connector":
                for k in BOX:
                    number(e.get(k), f"{eid}.{k}", k in {"w", "h"})
                require(e["x"] >= 0 and e["y"] >= 0 and e["x"] + e["w"] <= width + 1e-6
                        and e["y"] + e["h"] <= height + 1e-6, f"{sid}/{eid}: out of bounds")
            for k in ("color", "fill", "line"):
                if k in e:
                    color(e[k])
            if "size" in e:
                number(e["size"], "font size", True)
                if e["size"] < 14:
                    warnings.append(f"{sid}/{eid}: small type; review at intended presentation size")
            for k in ("bold", "header", "arrow"):
                if k in e:
                    require(type(e[k]) is bool, f"{k} must be boolean")
            if "align" in e:
                require(e["align"] in ("left", "center", "right"), "Unsupported alignment")
            if kind == "text" or "text" in e:
                string(e.get("text"), f"{eid}.text")
            if kind == "shape":
                require(e.get("shape", "rect") in SHAPES, "Unsupported shape")
            elif kind == "image":
                path = local_path(root, e.get("path"))
                string(e.get("alt"), "image alt description")
                require(isinstance(e.get("sha256"), str) and digest(path) == e["sha256"],
                        f"{eid}: source image hash mismatch")
                with Image.open(path) as im:
                    require(im.format in ("PNG", "JPEG"), "Use a reviewed local PNG/JPEG asset")
                    im.verify()
                with Image.open(path) as im:
                    im.load()
            elif kind == "table":
                rows = e.get("rows")
                require(isinstance(rows, list) and rows and isinstance(rows[0], list) and rows[0], "Empty table")
                require(all(isinstance(r, list) and len(r) == len(rows[0]) and
                            all(isinstance(c, str) for c in r) for r in rows), "Table must be rectangular strings")
                for row in rows:
                    for cell in row:
                        if cell:
                            string(cell, "table cell")
            elif kind == "chart":
                chart_type = e.get("chart_type")
                require(chart_type in CHARTS, "Unsupported chart type; use a task-local authoring extension")
                series = e.get("series")
                require(isinstance(series, list) and series, "Chart needs series")
                if chart_type != "scatter":
                    cats = e.get("categories")
                    require(isinstance(cats, list) and cats and all(isinstance(c, str) and c for c in cats),
                            "Category chart needs string categories")
                    require(len(cats) == len(set(cats)), "Duplicate chart categories")
                else:
                    require("categories" not in e, "Scatter requires numeric points, not categories")
                names = set()
                for s in series:
                    keys(s, {"name", "points"} if chart_type == "scatter" else {"name", "values"}, "series")
                    name = string(s.get("name"), "series name")
                    require(name not in names, "Duplicate series name")
                    names.add(name)
                    if chart_type == "scatter":
                        points = s.get("points")
                        require(isinstance(points, list) and points, "Scatter needs [x,y] pairs")
                        for p in points:
                            require(isinstance(p, list) and len(p) == 2, "Scatter needs [x,y] pairs")
                            for v in p:
                                number(v, "scatter coordinate")
                    else:
                        values = s.get("values")
                        require(isinstance(values, list) and len(values) == len(cats), "Chart data length mismatch")
                        for v in values:
                            number(v, "chart value")
                for k in ("x_title", "y_title"):
                    if k in e:
                        string(e[k], k)
                if not refs:
                    warnings.append(f"{sid}/{eid}: chart has no source reference; disclose synthetic data or supply provenance")
        for e in elements:
            if e["kind"] != "connector":
                continue
            require(e.get("routing", "straight") in ("straight", "elbow"), "Unsupported connector routing")
            for k in ("source", "target"):
                ref = e.get(k)
                keys(ref, {"id", "port"}, k)
                require(isinstance(ref.get("id"), str) and ref["id"] in by_id
                        and by_id[ref["id"]]["kind"] == "shape", "Connectors must attach to native shapes")
                require(ref.get("port") in PORTS, "Unknown connector port")
            require(e["source"]["id"] != e["target"]["id"], "Self-loops need an explicit custom route")
        content = [e for e in elements if e["kind"] in ("text", "table", "image", "chart") or e.get("text")]
        for i, a in enumerate(content):
            for b in content[i + 1:]:
                if (min(a["x"]+a["w"], b["x"]+b["w"]) - max(a["x"], b["x"]) > .02 and
                    min(a["y"]+a["h"], b["y"]+b["h"]) - max(a["y"], b["y"]) > .02):
                    warnings.append(f"{sid}: possible content overlap {a['id']} / {b['id']}; inspect rendering")
    warnings.append("Fonts are requested, not embedded or verified installed; inspect target-client rendering")
    return warnings


def load(path):
    path = path.resolve()
    spec = json.loads(path.read_text(encoding="utf-8"))
    warnings = validate(spec, path.parent)
    return path.parent, spec, warnings


def set_text(frame, value, e, theme, language):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.15
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[e.get("align", "left")]
        for run in p.runs:
            run.font.name = theme["font"]
            run.font.size = Pt(e.get("size", 22))
            run.font.bold = e.get("bold", False)
            run.font.color.rgb = color(e.get("color", theme["foreground"]))
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", language)
            ea = OxmlElement("a:ea"); ea.set("typeface", theme["cjk_font"]); rpr.append(ea)


def image_box(e, root):
    with Image.open(local_path(root, e["path"])) as im:
        factor = min(e["w"] / im.width, e["h"] / im.height)
        w, h = im.width * factor, im.height * factor
    return e["x"] + (e["w"]-w)/2, e["y"] + (e["h"]-h)/2, w, h


def make_chart(slide, e, box):
    if e["chart_type"] == "scatter":
        data = XyChartData()
        for series in e["series"]:
            s = data.add_series(series["name"])
            for x, y in series["points"]:
                s.add_data_point(x, y)
    else:
        data = CategoryChartData(); data.categories = e["categories"]
        for series in e["series"]:
            data.add_series(series["name"], series["values"])
    shape = slide.shapes.add_chart(CHARTS[e["chart_type"]], *box, data)
    chart = shape.chart
    chart.has_legend = len(e["series"]) > 1
    axes = (("x_title", chart.value_axis), ("y_title", chart.category_axis)) if e["chart_type"] == "bar" else (
        ("x_title", chart.category_axis), ("y_title", chart.value_axis))
    for field, axis in axes:
        if field in e:
            axis.has_title = True
            axis.axis_title.text_frame.text = e[field]
    return shape


def create(spec, root):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(spec["width"]), Inches(spec["height"])
    prs.core_properties.title = spec["title"]
    prs.core_properties.author = ""
    theme = spec["theme"]
    for page in spec["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(theme["background"])
        made = {}
        # Resolve all targets first; connectors are foreground objects, with true bindings.
        for e in page["elements"]:
            kind = e["kind"]
            if kind == "connector":
                continue
            box = [Inches(e[k]) for k in ("x", "y", "w", "h")]
            if kind == "text":
                sh = slide.shapes.add_textbox(*box)
                set_text(sh.text_frame, e["text"], e, theme, spec["language"])
            elif kind == "shape":
                sh = slide.shapes.add_shape(SHAPES[e.get("shape", "rect")], *box)
                sh.fill.solid(); sh.fill.fore_color.rgb = color(e.get("fill", theme["background"]))
                sh.line.color.rgb = color(e.get("line", theme["accent"]))
                if "text" in e:
                    set_text(sh.text_frame, e["text"], e, theme, spec["language"])
                    sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(.1)
                    sh.text_frame.margin_top = sh.text_frame.margin_bottom = Inches(.08)
            elif kind == "image":
                sh = slide.shapes.add_picture(str(local_path(root, e["path"])), *map(Inches, image_box(e, root)))
                sh._element.nvPicPr.cNvPr.set("descr", e["alt"])
            elif kind == "table":
                sh = slide.shapes.add_table(len(e["rows"]), len(e["rows"][0]), *box)
                sh.table.first_row = e.get("header", True)
                for r, row in enumerate(e["rows"]):
                    for c, value in enumerate(row):
                        cell = sh.table.cell(r, c)
                        cell.fill.solid(); cell.fill.fore_color.rgb = color(theme["background"])
                        opts = {"size": e.get("size", 18), "bold": r == 0 and e.get("header", True)}
                        set_text(cell.text_frame, value, opts, theme, spec["language"])
                        cell.margin_left = cell.margin_right = Inches(.06)
                        cell.margin_top = cell.margin_bottom = Inches(.04)
            elif kind == "chart":
                sh = make_chart(slide, e, box)
                sh.chart.font.name = theme["font"]
                sh.chart.font.size = Pt(16)
            sh.name = e["id"]
            made[e["id"]] = sh
        for e in page["elements"]:
            if e["kind"] != "connector":
                continue
            sh = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if e.get("routing") == "elbow" else MSO_CONNECTOR.STRAIGHT,
                                            0, 0, Inches(1), Inches(1))
            sh.begin_connect(made[e["source"]["id"]], PORTS[e["source"]["port"]])
            sh.end_connect(made[e["target"]["id"]], PORTS[e["target"]["port"]])
            sh.name = e["id"]
            sh.line.color.rgb = color(e.get("color", theme["accent"]))
            sh.line.width = Pt(1.5)
            if e.get("arrow", True):
                end = OxmlElement("a:tailEnd"); end.set("type", "triangle")
                sh.line._get_or_add_ln().append(end)
        if page.get("speaker_notes"):
            slide.notes_slide.notes_text_frame.text = page["speaker_notes"]
    return prs


def chart_values(series, tag):
    return [float(n.text) for n in series._element.xpath(f"./c:{tag}/c:numRef/c:numCache/c:pt/c:v")]


def inspect(prs, spec, root):
    """Inspect saved objects against the source; does not infer glyph layout."""
    require(len(prs.slides) == len(spec["slides"]), "Slide count mismatch")
    require(abs(prs.slide_width-Inches(spec["width"])) <= 2 and
            abs(prs.slide_height-Inches(spec["height"])) <= 2, "Canvas mismatch")
    stats = []
    for slide, page in zip(prs.slides, spec["slides"]):
        actual = {s.name: s for s in slide.shapes}
        require(len(actual) == len(slide.shapes) == len(page["elements"]), "Object count/name mismatch")
        counts = {}
        for e in page["elements"]:
            require(e["id"] in actual, f"Missing object {e['id']}")
            sh = actual[e["id"]]; kind = e["kind"]
            counts[kind] = counts.get(kind, 0) + 1
            if kind != "connector":
                expected = image_box(e, root) if kind == "image" else [e[k] for k in ("x", "y", "w", "h")]
                require(all(abs(a-Inches(b)) <= 3 for a, b in zip((sh.left, sh.top, sh.width, sh.height), expected)),
                        f"Geometry mismatch: {e['id']}")
            if kind in ("text", "shape"):
                require(sh.has_text_frame and sh.text == e.get("text", ""), f"Native text mismatch: {e['id']}")
                require(sh.shape_type == (MSO_SHAPE_TYPE.TEXT_BOX if kind == "text" else MSO_SHAPE_TYPE.AUTO_SHAPE),
                        f"Wrong native shape type: {e['id']}")
            elif kind == "image":
                require(sh.shape_type == MSO_SHAPE_TYPE.PICTURE and
                        hashlib.sha256(sh.image.blob).hexdigest() == e["sha256"], f"Image bytes changed: {e['id']}")
                require(all(getattr(sh, "crop_" + side) == 0 for side in ("top", "bottom", "left", "right")), "Unexpected image crop")
            elif kind == "table":
                require(sh.has_table, "Table is not native")
                require([[c.text for c in r.cells] for r in sh.table.rows] == e["rows"], "Table contents changed")
            elif kind == "chart":
                require(sh.has_chart and sh.chart.chart_type == CHARTS[e["chart_type"]], "Chart type mismatch")
                require(len(sh.chart.series) == len(e["series"]), "Chart series count mismatch")
                for actual_series, expected_series in zip(sh.chart.series, e["series"]):
                    require(actual_series.name == expected_series["name"], "Series name mismatch")
                    if e["chart_type"] == "scatter":
                        require(chart_values(actual_series, "xVal") == [p[0] for p in expected_series["points"]], "Scatter X values changed")
                        require(chart_values(actual_series, "yVal") == [p[1] for p in expected_series["points"]], "Scatter Y values changed")
                    else:
                        require(list(actual_series.values) == expected_series["values"], "Chart values changed")
                if e["chart_type"] != "scatter":
                    require([c.label for c in sh.chart.plots[0].categories] == e["categories"], "Chart categories changed")
            elif kind == "connector":
                require(sh.shape_type == MSO_SHAPE_TYPE.LINE, "Connector is not native")
                for key, tag in (("source", "stCxn"), ("target", "endCxn")):
                    ref = e[key]; nodes = sh._element.xpath(f".//a:{tag}")
                    require(len(nodes) == 1 and nodes[0].get("id") == str(actual[ref["id"]].shape_id)
                            and nodes[0].get("idx") == str(PORTS[ref["port"]]), "Connector is detached or attached to wrong port")
                arrow = sh._element.xpath(".//a:tailEnd")
                require((len(arrow) == 1 and arrow[0].get("type") == "triangle")
                        if e.get("arrow", True) else not arrow, "Arrow mismatch")
        if page.get("speaker_notes"):
            require(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text == page["speaker_notes"], "Notes mismatch")
        else:
            require(not slide.has_notes_slide, "Unexpected speaker notes")
        stats.append({"id": page["id"], "objects": counts})
    return stats


def output_paths(root, name):
    path = local_path(root, name)
    require(path.suffix.lower() == ".pptx", "Output must have .pptx extension")
    return path, path.with_suffix(".build.json")


def source_assets(spec, root):
    return {e["path"]: digest(local_path(root, e["path"])) for s in spec["slides"]
            for e in s["elements"] if e["kind"] == "image"}


def build(manifest, name):
    root, spec, warnings = load(manifest)
    output, receipt = output_paths(root, name)
    require(not output.exists() and not receipt.exists(), "Output/release already exists; choose a new revision name")
    prs = create(spec, root)
    buffer = BytesIO(); prs.save(buffer); data = buffer.getvalue()
    stats = inspect(Presentation(BytesIO(data)), spec, root)
    report = {"schema_version": 1, "source_sha256": digest(manifest),
              "pptx_sha256": hashlib.sha256(data).hexdigest(), "assets": source_assets(spec, root),
              "slides": stats, "structural_check": "passed", "content_review": "not_performed",
              "native_render": "not_performed", "visual_review": "not_performed", "warnings": warnings}
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects previous releases, even when another build races.
    created = []
    try:
        with output.open("xb") as f:
            created.append(output); f.write(data)
        with receipt.open("x", encoding="utf-8") as f:
            created.append(receipt); json.dump(report, f, ensure_ascii=False, indent=2); f.write("\n")
    except OSError:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return {"pptx": str(output), "receipt": str(receipt), **report}


def check(manifest, name):
    root, spec, warnings = load(manifest)
    output, receipt = output_paths(root, name)
    report = json.loads(receipt.read_text(encoding="utf-8"))
    require(report.get("source_sha256") == digest(manifest), "Stale source: rebuild a new revision")
    require(report.get("pptx_sha256") == digest(output), "PPTX changed since build; reconcile manual edits before rebuilding")
    require(report.get("assets") == source_assets(spec, root), "Source assets changed")
    stats = inspect(Presentation(str(output)), spec, root)
    return {"pptx": str(output), "slides": stats, "structural_check": "passed", "warnings": warnings,
            "content_review": "not_performed", "native_render": "not_performed", "visual_review": "not_performed"}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("build", "check"))
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--output", required=True, help="New PPTX path relative to manifest directory")
    args = parser.parse_args(argv)
    try:
        result = (build if args.command == "build" else check)(args.manifest, args.output)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, TypeError, KeyError) as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
